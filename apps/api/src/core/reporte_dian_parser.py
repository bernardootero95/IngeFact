"""Lectura del reporte de documentos que descarga el tenant del portal de la
DIAN (Facturando Electronicamente > Documentos > Exportar Excel).

Se usa en la carga masiva de Facturas Recibidas: del reporte solo salen los
CUFE candidatos; los datos reales de cada factura se vuelven a leer de la
DIAN (via Alegra) antes de registrarla. Funcion pura, sin BD ni red.

Las columnas se reconocen por su nombre (sin importar tildes, mayusculas ni
orden), para no depender del orden en que el portal las exporte.
"""

import re
import unicodedata
import zipfile
from dataclasses import dataclass
from io import BytesIO

from openpyxl import load_workbook

MAX_FILAS = 5000

# Encabezado normalizado -> campo de FilaReporteDian.
COLUMNAS = {
    "tipo de documento": "tipo_documento",
    "cufe/cude": "cufe",
    "folio": "folio",
    "prefijo": "prefijo",
    "forma de pago": "forma_pago",
    "fecha emision": "fecha_emision",
    "nit emisor": "nit_emisor",
    "nombre emisor": "nombre_emisor",
    "nit receptor": "nit_receptor",
    "total": "total",
    "estado": "estado",
    "grupo": "grupo",
}
COLUMNAS_OBLIGATORIAS = ("tipo_documento", "cufe", "nit_receptor", "forma_pago", "grupo")


class ReporteDianInvalido(ValueError):
    """El archivo no es un reporte de documentos de la DIAN legible."""


@dataclass(frozen=True)
class FilaReporteDian:
    tipo_documento: str
    cufe: str
    folio: str
    prefijo: str
    forma_pago: str
    fecha_emision: str
    nit_emisor: str
    nombre_emisor: str
    nit_receptor: str
    total: float | None
    estado: str
    grupo: str


def normalizar_texto(valor) -> str:
    """Minusculas, sin tildes y sin espacios repetidos."""
    texto = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", texto).strip().lower()


def _encabezado(valor) -> str | None:
    return COLUMNAS.get(normalizar_texto(valor))


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _numero(valor) -> float | None:
    if valor is None or valor == "":
        return None
    try:
        return float(valor)
    except (TypeError, ValueError):
        return None


def parsear_reporte_dian(contenido: bytes) -> list[FilaReporteDian]:
    try:
        libro = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise ReporteDianInvalido("El archivo no es un Excel válido. Descarga de nuevo el reporte de la DIAN.") from exc

    try:
        filas = libro.worksheets[0].iter_rows(values_only=True)
        encabezados = next(filas, None)
        if not encabezados:
            raise ReporteDianInvalido("El archivo está vacío.")
        posiciones = {campo: i for i, valor in enumerate(encabezados) if (campo := _encabezado(valor))}
        faltantes = [campo for campo in COLUMNAS_OBLIGATORIAS if campo not in posiciones]
        if faltantes:
            raise ReporteDianInvalido(
                "Este archivo no parece el reporte de documentos de la DIAN: faltan columnas como "
                "\"Tipo de documento\", \"CUFE/CUDE\", \"NIT Receptor\", \"Forma de Pago\" o \"Grupo\"."
            )

        resultado: list[FilaReporteDian] = []
        for fila in filas:
            if not any(v not in (None, "") for v in fila):
                continue
            if len(resultado) >= MAX_FILAS:
                raise ReporteDianInvalido(f"El reporte tiene más de {MAX_FILAS} filas. Descárgalo por un rango de fechas más corto.")

            def valor(campo: str):
                i = posiciones.get(campo)
                return fila[i] if i is not None and i < len(fila) else None

            resultado.append(
                FilaReporteDian(
                    tipo_documento=_texto(valor("tipo_documento")),
                    cufe=_texto(valor("cufe")),
                    folio=_texto(valor("folio")),
                    prefijo=_texto(valor("prefijo")),
                    forma_pago=_texto(valor("forma_pago")),
                    fecha_emision=_texto(valor("fecha_emision")),
                    nit_emisor=_texto(valor("nit_emisor")),
                    nombre_emisor=_texto(valor("nombre_emisor")),
                    nit_receptor=_texto(valor("nit_receptor")),
                    total=_numero(valor("total")),
                    estado=_texto(valor("estado")),
                    grupo=_texto(valor("grupo")),
                )
            )
        return resultado
    finally:
        libro.close()
